from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from .case_deidentifier import deidentify_text, looks_like_personal_name, sanitize_label


SUPPORTED_EXTS = {".doc", ".docx", ".pdf", ".xls", ".xlsx"}
_NON_CASE_NAME = re.compile(r"规范|范文|书写|模板|要求|送：")
_SECTION_CUT = re.compile(r"[\n\r]|现病史|既往史|个人史|家族史|婚育史|月经史|体格检查|辅助检查|诊疗经过|过敏史|初步诊断|出院诊断")


class _DocExtractor:
    """Reuse one hidden Word instance across all legacy .doc files."""

    def __init__(self) -> None:
        self._word: Any = None

    def _ensure_word(self) -> Any:
        if self._word is None:
            import win32com.client

            self._word = win32com.client.DispatchEx("Word.Application")
            self._word.Visible = False
            self._word.DisplayAlerts = 0
        return self._word

    def extract_doc(self, path: Path) -> str:
        word = self._ensure_word()
        document = word.Documents.Open(str(path), ReadOnly=True, AddToRecentFiles=False, Visible=False)
        try:
            return (document.Content.Text or "").strip()
        finally:
            document.Close(False)

    def close(self) -> None:
        if self._word is None:
            return
        try:
            self._word.Quit()
        except Exception:
            pass
        finally:
            self._word = None


def extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs)


def extract_pdf(path: Path) -> str:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]
    return "\n".join(page.strip() for page in pages if page.strip())


def extract_xlsx(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    rows: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = [str(cell).strip() for cell in row if cell not in (None, "")]
            if values:
                rows.append(" | ".join(values))
    workbook.close()
    return "\n".join(rows)


def extract_xls(path: Path) -> str:
    import win32com.client

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = 0
    try:
        workbook = excel.Workbooks.Open(str(path), ReadOnly=True)
        try:
            lines: list[str] = []
            for sheet in workbook.Worksheets:
                used = sheet.UsedRange
                if used is None:
                    continue
                for row in used.Value or []:
                    values = [str(cell).strip() for cell in row if cell not in (None, "")]
                    if values:
                        lines.append(" | ".join(values))
            return "\n".join(lines)
        finally:
            workbook.Close(False)
    finally:
        excel.Quit()


def extract_file(path: Path, doc_extractor: _DocExtractor) -> str:
    suffix = path.suffix.lower()
    if suffix == ".doc":
        return doc_extractor.extract_doc(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".xlsx":
        return extract_xlsx(path)
    if suffix == ".xls":
        return extract_xls(path)
    raise ValueError(f"不支持的文件类型：{suffix}")


def _case_files(folder: Path) -> list[Path]:
    files = [
        path
        for path in folder.iterdir()
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTS and not _NON_CASE_NAME.search(path.name)
    ]
    return sorted(files, key=lambda item: item.name)


def _pick_main_file(files: list[Path]) -> Path:
    def score(path: Path) -> int:
        name = path.name
        value = 0
        if re.search(r"住院病历|大病历|大病例|病历", name):
            value += 100
        if re.search(r"首记|首次病程|首程|出院记录", name):
            value += 40
        if re.search(r"病程", name):
            value += 20
        value += min(path.stat().st_size // 1024, 60)
        return value

    return max(files, key=score)


def _extract_section(text: str, key: str, max_length: int) -> str:
    match = re.search(rf"{key}\s*[:：]?\s*(.{{0,{max_length}}})", text)
    if not match:
        return ""
    value = _SECTION_CUT.split(match.group(1), maxsplit=1)[0]
    return value.strip(" \u3000:：,，。;；\n\r")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _build_entry(category: str, folder: Path, text: str) -> dict[str, Any] | None:
    masked = deidentify_text(text[:12000])
    if len(masked.text.strip()) < 40:
        return None

    label = sanitize_label(folder.name)
    parent_label = sanitize_label(folder.parent.name) if folder.parent else ""
    if not label or looks_like_personal_name(label):
        if parent_label and parent_label != category and not looks_like_personal_name(parent_label):
            label = parent_label
        else:
            label = "未命名病例"

    chief_complaint = _extract_section(masked.text, "主诉", 140)
    present_illness = _extract_section(masked.text, "现病史", 900)
    diagnosis = _extract_section(masked.text, r"(出院诊断|最后诊断|初步诊断|入院诊断|诊断)", 140)
    if not diagnosis or diagnosis in {"诊断", "待查", "待确诊", "无"} or len(diagnosis) <= 2:
        diagnosis = label
    title = label if label.endswith(("病例", "病历", "案例")) else f"{label}病例"

    entry_id = f"ckb_{uuid4().hex[:10]}"
    timestamp = _now()
    return {
        "id": entry_id,
        "title": title,
        "category": category,
        "diagnosis": diagnosis,
        "chief_complaint": chief_complaint,
        "present_illness": present_illness,
        "content": masked.text[:10000],
        "source": f"{category}/{label}" if label else category,
        "status": "已脱敏入库",
        "anonymized": True,
        "imported": True,
        "pii_removed": masked.pii_types,
        "risk_flags": masked.risk_flags,
        "created_at": timestamp,
        "updated_at": timestamp,
    }


def import_case_library(source_dir: str | Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    root = Path(source_dir)
    if not root.exists():
        raise FileNotFoundError(f"素材目录不存在：{root}")

    entries: list[dict[str, Any]] = []
    scanned_files = 0
    skipped = 0
    failed = 0
    categories: dict[str, int] = {}
    failures: list[str] = []
    doc_extractor = _DocExtractor()

    try:
        for category_dir in sorted([path for path in root.iterdir() if path.is_dir()]):
            category = category_dir.name
            folders = [category_dir]
            folders.extend(sorted([path for path in category_dir.rglob("*") if path.is_dir()]))
            for folder in folders:
                files = _case_files(folder)
                if not files:
                    continue
                scanned_files += len(files)
                text = ""
                for candidate in [_pick_main_file(files), *[item for item in files if item != _pick_main_file(files)]]:
                    try:
                        text = extract_file(candidate, doc_extractor)
                        if text.strip():
                            break
                    except Exception as exc:
                        failures.append(f"{candidate.name}: {exc}")
                if not text.strip():
                    skipped += 1
                    continue
                entry = _build_entry(category, folder, text)
                if entry is None:
                    skipped += 1
                    continue
                entries.append(entry)
                categories[category] = categories.get(category, 0) + 1
    finally:
        doc_extractor.close()

    report = {
        "source_dir": str(root),
        "scanned_files": scanned_files,
        "imported": len(entries),
        "skipped": skipped,
        "failed": failed,
        "categories": categories,
        "failures": failures[:20],
    }
    return entries, report
